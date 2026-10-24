from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from platforms.column_mappings import save_column_mapping
from platforms.xlsx_reports import (
    MONTHLY_REPORT,
    QUARTERLY_REPORT,
    TRANSACTION_REPORT,
    XlsxReportError,
    analyze_report_file,
    detect_platform_key,
    load_report_rows,
    load_workbook_preview,
    match_rows_to_products,
    report_platform_key,
    resolve_report_column_names,
    rows_for_monthly_summary,
    scan_report_directory,
)


def _save_mapping_for_table(
    report_path,
    database_path,
    report_type,
    columns,
    sheet_name=None,
    header_row=1,
):
    workbook = load_workbook(report_path, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        headers = [
            sheet.cell(row=header_row, column=column).value
            for column in range(1, (sheet.max_column or 1) + 1)
        ]
    finally:
        workbook.close()
    save_column_mapping(
        report_platform_key(report_path),
        headers,
        report_type,
        columns,
        database_path=database_path,
    )


@pytest.fixture
def mapping_database(tmp_path):
    return tmp_path / "mappings.sqlite"


def _write_legimi_report(path):
    """Raport Legimi: naglowek w 8. wierszu i wiersz 'Razem' na koncu."""
    workbook = Workbook()
    sheet = workbook.active
    sheet["D3"] = "Lista transakcji"
    sheet["D5"] = "Od: 01.01.2026 Do: 13.08.2026"
    sheet.append([])

    header_row = 8
    headers = [
        "Tytuł", "", "ISBN", "", "Format", "Data sprzedaży",
        "Typ sprzedaży", "Wynagrodzenie netto", "Wynagrodzenie brutto",
    ]
    for column, value in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=value)

    data = [
        ["Lustro Iluzji", "", "9788397514638", "", "audiobook", "2026-08-08",
         "Katalog bez limitu", 23.33, 24.5],
        ["Lustro Iluzji", "", "9788397514638", "", "audiobook", "2026-08-04",
         "Katalog bez limitu", 23.33, 24.5],
        ["Życiowa Rozgrywka", "", "9788368950052", "", "audiobook", "2026-08-11",
         "Katalog bez limitu", 20.0, 21.0],
        ["Razem", "", "", "", "", "", "", 66.66, 70.0],
    ]
    for offset, row_values in enumerate(data, start=header_row + 1):
        for column, value in enumerate(row_values, start=1):
            sheet.cell(row=offset, column=column, value=value)

    workbook.save(path)
    return path


def _write_storytel_report(path):
    """Raport Storytel: arkusz z faktura oraz arkusz z ksiazkami."""
    workbook = Workbook()
    report_sheet = workbook.active
    report_sheet.title = "Report"
    report_sheet["B3"] = "Report"
    report_sheet["E12"] = "Report number:"
    report_sheet["F12"] = "27530001"

    books = workbook.create_sheet("A-books")
    books.append([
        "Author(s)", "Audio Book Title", "ISBN", "Pool", "Country",
        "Price model", "Quantity/no. of units", "Net receipts per hour",
        "ECB exchange rate", "Net receipts per hour (PLN)",
        "Book length (in hours)", "Price per Unit (PLN)", "Remuneration (PLN)",
        "VAT (%)", "Publisher", "Imprint", "Consumption dates",
    ])
    books.append([
        "Karolina Bast Kolasa", "Pamiętniki Kaletnickie", "9788398141550",
        "PL Primary Pool", "pl", "Revenue share 40%/hour", 1.0, 0.79, "1 PLN = 1 PLN",
        0.79, 1.98, 0.63, 100.0, 5.0, "Audio Opus", "Audio Opus",
        "2026-06-30 to 2026-06-30",
    ])

    workbook.save(path)
    return path


@pytest.fixture
def legimi_report(tmp_path):
    return _write_legimi_report(tmp_path / "Legimi.xlsx")


@pytest.fixture
def storytel_report(tmp_path):
    return _write_storytel_report(tmp_path / "Storytel.xlsx")


@pytest.fixture
def products():
    return [
        {"id": 1, "title": "Lustro iluzji", "isbn": "978-83-9751-463-8"},
        {"id": 2, "title": "Życiowa Rozgrywka", "isbn": "978-83-68950-05-2"},
        # Celowo inny ISBN niz w raporcie - dopasowanie ma pojsc po tytule.
        {"id": 3, "title": "Pamiętniki Kaletnickie", "isbn": "978-83-68950-04-5"},
    ]


class TestDetectPlatform:

    def test_detects_platform_from_file_name(self, tmp_path):
        assert detect_platform_key(tmp_path / "Legimi.xlsx") == "legimi"
        assert detect_platform_key(tmp_path / "Storytel.xlsx") == "storytel"

    def test_returns_none_for_unknown_file_name(self, tmp_path):
        assert detect_platform_key(tmp_path / "Sprzedaz.xlsx") is None


class TestLoadReportRows:

    def test_reads_legimi_transactions(self, legimi_report, mapping_database):
        _save_mapping_for_table(
            legimi_report,
            mapping_database,
            TRANSACTION_REPORT,
            {
                "title": 0,
                "isbn": 2,
                "transaction_date": 5,
                "sale_type": 6,
                "net": 7,
                "gross": 8,
            },
            header_row=8,
        )
        platform_key, rows = load_report_rows(
            legimi_report,
            mapping_database_path=mapping_database,
        )

        assert platform_key == "legimi"
        # Wiersz "Razem" nie jest pozycja raportu.
        assert len(rows) == 3
        assert rows[0].title == "Lustro Iluzji"
        assert rows[0].isbn == "9788397514638"
        assert rows[0].net == 23.33
        assert rows[0].gross == 24.5
        assert rows[0].month == "2026-08"
        assert rows[0].report_type == TRANSACTION_REPORT

    def test_reads_storytel_books_sheet_and_computes_gross_from_vat(
        self,
        storytel_report,
        mapping_database,
    ):
        _save_mapping_for_table(
            storytel_report,
            mapping_database,
            MONTHLY_REPORT,
            {"title": 1, "isbn": 2, "net": 12, "vat": 13, "period": 16},
            sheet_name="A-books",
        )
        platform_key, rows = load_report_rows(
            storytel_report,
            mapping_database_path=mapping_database,
        )

        assert platform_key == "storytel"
        assert len(rows) == 1
        assert rows[0].title == "Pamiętniki Kaletnickie"
        assert rows[0].net == 100.0
        assert rows[0].gross == pytest.approx(105.0)
        assert rows[0].month == "2026-06"
        assert rows[0].report_type == MONTHLY_REPORT

    def test_reads_irregular_storytel_range_as_quarterly_settlement(
        self,
        storytel_report,
        mapping_database,
    ):
        _save_mapping_for_table(
            storytel_report,
            mapping_database,
            QUARTERLY_REPORT,
            {
                "title": 1,
                "isbn": 2,
                "quantity": 6,
                "net": 12,
                "vat": 13,
                "period": 16,
            },
            sheet_name="A-books",
        )

        platform_key, rows = load_report_rows(
            storytel_report,
            mapping_database_path=mapping_database,
        )

        assert platform_key == "storytel"
        assert len(rows) == 1
        assert rows[0].report_type == QUARTERLY_REPORT
        assert rows[0].period_start == "2026-06-30"
        assert rows[0].period_end == "2026-06-30"
        assert rows[0].quantity == 1.0
        assert rows_for_monthly_summary(rows) == []

    def test_uses_mapping_when_file_name_is_unknown(self, tmp_path, mapping_database):
        report_path = _write_legimi_report(tmp_path / "raport-2026.xlsx")
        _save_mapping_for_table(
            report_path,
            mapping_database,
            TRANSACTION_REPORT,
            {"title": 0, "isbn": 2, "transaction_date": 5, "net": 7},
            header_row=8,
        )

        platform_key, rows = load_report_rows(
            report_path,
            mapping_database_path=mapping_database,
        )

        assert platform_key == "raport-2026"
        assert len(rows) == 3

    def test_raises_when_file_has_no_report_table(self, tmp_path):
        workbook = Workbook()
        workbook.active["A1"] = "Notatki"
        report_path = tmp_path / "Legimi.xlsx"
        workbook.save(report_path)

        with pytest.raises(XlsxReportError):
            load_report_rows(report_path)

    def test_raises_for_missing_file(self, tmp_path):
        with pytest.raises(XlsxReportError):
            load_report_rows(tmp_path / "Brak.xlsx")


class TestMatchRowsToProducts:

    def test_sums_rows_of_the_same_product(
        self,
        legimi_report,
        products,
        mapping_database,
    ):
        _save_mapping_for_table(
            legimi_report,
            mapping_database,
            TRANSACTION_REPORT,
            {
                "title": 0,
                "isbn": 2,
                "transaction_date": 5,
                "net": 7,
                "gross": 8,
            },
            header_row=8,
        )
        _, rows = load_report_rows(
            legimi_report,
            mapping_database_path=mapping_database,
        )

        matched, unmatched = match_rows_to_products(rows, products)

        assert unmatched == []
        by_id = {summary.product["id"]: summary for summary in matched}
        assert by_id[1].net == 46.66
        assert by_id[1].gross == 49.0
        assert by_id[1].rows == 2
        assert by_id[2].net == 20.0

    def test_matches_by_title_when_isbn_differs(
        self,
        storytel_report,
        products,
        mapping_database,
    ):
        _save_mapping_for_table(
            storytel_report,
            mapping_database,
            MONTHLY_REPORT,
            {"title": 1, "isbn": 2, "net": 12, "period": 16},
            sheet_name="A-books",
        )
        _, rows = load_report_rows(
            storytel_report,
            mapping_database_path=mapping_database,
        )

        matched, unmatched = match_rows_to_products(rows, products)

        assert unmatched == []
        assert len(matched) == 1
        assert matched[0].product["id"] == 3
        assert matched[0].net == 100.0

    def test_reports_titles_missing_in_database(
        self,
        legimi_report,
        mapping_database,
    ):
        _save_mapping_for_table(
            legimi_report,
            mapping_database,
            TRANSACTION_REPORT,
            {
                "title": 0,
                "isbn": 2,
                "transaction_date": 5,
                "net": 7,
                "gross": 8,
            },
            header_row=8,
        )
        _, rows = load_report_rows(
            legimi_report,
            mapping_database_path=mapping_database,
        )

        matched, unmatched = match_rows_to_products(rows, [])

        assert matched == []
        titles = {entry.title for entry in unmatched}
        assert titles == {"Lustro Iluzji", "Życiowa Rozgrywka"}
        by_title = {entry.title: entry for entry in unmatched}
        assert by_title["Lustro Iluzji"].net == 46.66
        assert by_title["Lustro Iluzji"].gross == 49.0
        assert by_title["Lustro Iluzji"].rows == 2

    def test_gross_stays_empty_when_report_has_no_gross_column(
        self,
        tmp_path,
        products,
        mapping_database,
    ):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Tytuł", "ISBN", "Okres", "Wynagrodzenie netto"])
        sheet.append(["Lustro iluzji", "9788397514638", "2026-08", 10.0])
        report_path = tmp_path / "Legimi.xlsx"
        workbook.save(report_path)
        _save_mapping_for_table(
            report_path,
            mapping_database,
            MONTHLY_REPORT,
            {"title": 0, "isbn": 1, "period": 2, "net": 3},
        )

        _, rows = load_report_rows(
            report_path,
            mapping_database_path=mapping_database,
        )
        matched, _ = match_rows_to_products(rows, products)

        assert rows[0].gross is None
        assert matched[0].gross is None


def _write_audioteka_report(path):
    """Raport Audioteki: naglowek w 1. wierszu i wiersz 'Suma:' na koncu."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "Okres", "Tytuł", "Długość audiobooka", "Minuty", "Odsłuchy",
        "Stawka", "Kwota należna", "Waluta", "ISBN",
    ])
    sheet.append(["2025-05", "Duch Gór", 18.0, 72.0, 4.03, 0.25, 1.01, "PLN", ""])
    sheet.append([
        "2025-05", "Zasady gry", 455.0, 1254.0, 2.76, 6.37, 17.58, "PLN",
        "978-83-68244-31-1",
    ])
    sheet.append(["Suma:", "", "", "", "", "", 18.59, "", ""])
    sheet.append(["", "", "", "", "", "", "Kwota nie zawiera podatku VAT.", "", ""])

    workbook.save(path)
    return path


@pytest.fixture
def audioteka_report(tmp_path):
    return _write_audioteka_report(tmp_path / "Audioteka 2025-05 - Stowarzyszenie.xlsx")


class TestAudiotekaReports:

    def test_detects_platform_from_file_name(self, audioteka_report):
        assert detect_platform_key(audioteka_report) == "audioteka"

    def test_reads_rows_and_skips_summary_row(
        self,
        audioteka_report,
        mapping_database,
    ):
        _save_mapping_for_table(
            audioteka_report,
            mapping_database,
            MONTHLY_REPORT,
            {"period": 0, "title": 1, "net": 6, "isbn": 8},
        )
        platform_key, rows = load_report_rows(
            audioteka_report,
            mapping_database_path=mapping_database,
        )

        assert platform_key == "audioteka"
        assert len(rows) == 1
        assert [row.title for row in rows] == ["Zasady gry"]
        assert rows[0].net == 17.58
        assert rows[0].isbn == "978-83-68244-31-1"
        assert rows[0].month == "2025-05"
        # Audioteka nie podaje kwoty brutto ani stawki VAT.
        assert rows[0].gross is None

    def test_matches_by_isbn_written_with_dashes(
        self,
        audioteka_report,
        mapping_database,
    ):
        _save_mapping_for_table(
            audioteka_report,
            mapping_database,
            MONTHLY_REPORT,
            {"period": 0, "title": 1, "net": 6, "isbn": 8},
        )
        _, rows = load_report_rows(
            audioteka_report,
            mapping_database_path=mapping_database,
        )
        products = [{"id": 8, "title": "Inny tytuł", "isbn": "9788368244311"}]

        matched, unmatched = match_rows_to_products(rows, products)

        assert len(matched) == 1
        assert matched[0].product["id"] == 8
        assert matched[0].net == 17.58
        assert unmatched == []


class TestGenericReportLayouts:

    def test_resolves_only_exact_configured_names(self, tmp_path):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Podsumowanie raportu"
        sheet.append([])
        sheet.append(["Nazwa utworu", "Kod produktu", "Miesiąc", "Kwota"])
        sheet.append(["Książka", "9781234567890", "2026-07", 10])
        report_path = tmp_path / "Platforma.xlsx"
        workbook.save(report_path)

        match = resolve_report_column_names(
            report_path,
            {
                "title": "Nazwa utworu",
                "isbn": "Kod produktu",
                "period": "Miesiąc",
                "net": "Kwota",
            },
        )

        assert match.sheet_name == "Sheet"
        assert match.header_row == 3
        assert match.columns == {"title": 0, "isbn": 1, "period": 2, "net": 3}

    def test_workbook_preview_exposes_cells_for_picker(self, tmp_path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Raport"
        sheet.append(["Nazwa utworu", "Kod produktu"])
        sheet.append(["Książka", "9781234567890"])
        report_path = tmp_path / "Platforma.xlsx"
        workbook.save(report_path)

        previews = load_workbook_preview(report_path)

        assert len(previews) == 1
        assert previews[0].name == "Raport"
        assert previews[0].rows[0] == ("Nazwa utworu", "Kod produktu")
        assert previews[0].truncated is False

    def test_workbook_preview_includes_all_filled_rows_by_default(self, tmp_path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Nagłówek"])
        for row_number in range(1, 526):
            sheet.append([f"Pozycja {row_number}"])
        report_path = tmp_path / "Długi raport.xlsx"
        workbook.save(report_path)

        preview = load_workbook_preview(report_path)[0]

        assert preview.total_rows == 526
        assert len(preview.rows) == 526
        assert preview.rows[-1] == ("Pozycja 525",)
        assert preview.truncated is False

    def test_workbook_preview_reads_sheet_without_dimension_metadata(self, tmp_path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Channel", "Royalty Payment Currency", "Royalty Payment"])
        sheet.append(["No transactions during this period"])
        report_path = tmp_path / "Voices.xlsx"
        workbook.save(report_path)

        rewritten_path = tmp_path / "Voices-without-dimension.xlsx"
        with ZipFile(report_path, "r") as source, ZipFile(
            rewritten_path,
            "w",
            ZIP_DEFLATED,
        ) as target:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    start = content.find(b"<dimension ")
                    if start >= 0:
                        end = content.find(b"/>", start)
                        content = content[:start] + content[end + 2 :]
                target.writestr(entry, content)

        preview = load_workbook_preview(rewritten_path)[0]

        assert preview.total_rows == 2
        assert preview.total_columns == 3
        assert preview.rows[0] == (
            "Channel",
            "Royalty Payment Currency",
            "Royalty Payment",
        )
        assert preview.rows[1] == (
            "No transactions during this period",
            "",
            "",
        )

    def test_workbook_preview_preserves_empty_cells_and_formula_text(self, tmp_path):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Tytuł"
        sheet["C3"] = "=1+1"
        report_path = tmp_path / "Układ.xlsx"
        workbook.save(report_path)

        preview = load_workbook_preview(report_path)[0]

        assert preview.total_rows == 3
        assert preview.total_columns == 3
        assert preview.rows == (
            ("Tytuł", "", ""),
            ("", "", ""),
            ("", "", "=1+1"),
        )

    def test_remembers_unknown_columns_for_platform_and_layout(self, tmp_path):
        workbook = Workbook()
        sheet = workbook.active
        headers = ["Work Name", "EAN Code", "Statement Cycle", "Payable To Us"]
        sheet.append(headers)
        sheet.append(["Nowa książka", "9781234567890", "2026-07", "19,99"])
        report_path = tmp_path / "Nowa platforma.xlsx"
        workbook.save(report_path)
        mapping_database = tmp_path / "mappings.sqlite"

        with pytest.raises(XlsxReportError):
            load_report_rows(
                report_path,
                mapping_database_path=mapping_database,
            )

        match = resolve_report_column_names(
            report_path,
            {
                "title": "Work Name",
                "isbn": "EAN Code",
                "period": "Statement Cycle",
                "net": "Payable To Us",
            },
        )
        save_column_mapping(
            "nowa platforma",
            match.headers,
            MONTHLY_REPORT,
            match.columns,
            database_path=mapping_database,
        )

        platform_key, rows = load_report_rows(
            report_path,
            mapping_database_path=mapping_database,
        )

        assert platform_key == "nowa platforma"
        assert len(rows) == 1
        assert rows[0].title == "Nowa książka"
        assert rows[0].isbn == "9781234567890"
        assert rows[0].month == "2026-07"
        assert rows[0].net == 19.99
        analysis = analyze_report_file(
            report_path,
            expected_year=2026,
            mapping_database_path=mapping_database,
        )
        assert analysis.recognized is True

    def test_saved_layout_does_not_affect_another_platform(self, tmp_path):
        headers = ["Work Name", "EAN Code", "Statement Cycle", "Payable To Us"]
        first_path = tmp_path / "Platforma A.xlsx"
        second_path = tmp_path / "Platforma B.xlsx"
        for report_path in (first_path, second_path):
            workbook = Workbook()
            workbook.active.append(headers)
            workbook.active.append(["Książka", "9781234567890", "2026-07", 10])
            workbook.save(report_path)
        mapping_database = tmp_path / "mappings.sqlite"
        match = resolve_report_column_names(
            first_path,
            {
                "title": "Work Name",
                "isbn": "EAN Code",
                "period": "Statement Cycle",
                "net": "Payable To Us",
            },
        )
        save_column_mapping(
            "platforma a",
            match.headers,
            MONTHLY_REPORT,
            match.columns,
            database_path=mapping_database,
        )

        _, rows = load_report_rows(
            first_path,
            mapping_database_path=mapping_database,
        )
        assert len(rows) == 1
        with pytest.raises(XlsxReportError):
            load_report_rows(
                second_path,
                mapping_database_path=mapping_database,
            )

    def test_does_not_use_builtin_column_aliases(self, tmp_path, mapping_database):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            ["Book Name", "Digital ISBN", "Reporting Month", "Net Earnings"]
        )
        sheet.append(["Nowa książka", 9788397514638, "2026-07", "1,234.56 USD"])
        report_path = tmp_path / "Nowa platforma.xlsx"
        workbook.save(report_path)

        with pytest.raises(XlsxReportError, match="brak zapisanego mapowania"):
            load_report_rows(
                report_path,
                mapping_database_path=mapping_database,
            )

    def test_searches_entire_sheet_and_reads_multiple_mapped_tables(
        self,
        tmp_path,
        mapping_database,
    ):
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(row=45, column=1, value="Title")
        sheet.cell(row=45, column=2, value="ISBN")
        sheet.cell(row=45, column=3, value="Transaction Date")
        sheet.cell(row=45, column=4, value="Net Royalty")
        sheet.append(["Pierwsza", "978-1", "2026-01-12", 10])
        sheet.cell(row=50, column=1, value="Book Title")
        sheet.cell(row=50, column=2, value="ISBN-13")
        sheet.cell(row=50, column=3, value="Period")
        sheet.cell(row=50, column=4, value="Amount Due")
        sheet.cell(row=51, column=1, value="Druga")
        sheet.cell(row=51, column=2, value="978-2")
        sheet.cell(row=51, column=3, value="2026-02")
        sheet.cell(row=51, column=4, value=20)
        report_path = tmp_path / "Dwie tabele.xlsx"
        workbook.save(report_path)
        _save_mapping_for_table(
            report_path,
            mapping_database,
            TRANSACTION_REPORT,
            {"title": 0, "isbn": 1, "transaction_date": 2, "net": 3},
            header_row=45,
        )
        _save_mapping_for_table(
            report_path,
            mapping_database,
            MONTHLY_REPORT,
            {"title": 0, "isbn": 1, "period": 2, "net": 3},
            header_row=50,
        )

        _, rows = load_report_rows(
            report_path,
            mapping_database_path=mapping_database,
        )

        assert [(row.title, row.report_type) for row in rows] == [
            ("Pierwsza", TRANSACTION_REPORT),
            ("Druga", MONTHLY_REPORT),
        ]

    def test_analysis_marks_recognized_and_unrecognized_files(
        self,
        tmp_path,
        mapping_database,
    ):
        reports_root = tmp_path / "Raporty Platform"
        year_directory = reports_root / "2026"
        year_directory.mkdir(parents=True)
        recognized_path = _write_legimi_report(year_directory / "Rozpoznany.xlsx")
        unrecognized_path = year_directory / "Nieznany.xlsx"
        workbook = Workbook()
        workbook.active.append(["Notatki", "Bez tabeli"])
        workbook.save(unrecognized_path)
        unsupported_path = year_directory / "Załącznik.pdf"
        unsupported_path.write_bytes(b"not a report")
        (year_directory / "pomocniczy.json").write_text("{}", encoding="utf-8")
        trashed_path = reports_root / ".Kosz" / "batch" / "2026" / "Kosz.xlsx"
        trashed_path.parent.mkdir(parents=True)
        trashed_path.write_bytes(b"deleted report")
        _save_mapping_for_table(
            recognized_path,
            mapping_database,
            TRANSACTION_REPORT,
            {"title": 0, "isbn": 2, "transaction_date": 5, "net": 7},
            header_row=8,
        )

        analyses = scan_report_directory(
            reports_root,
            mapping_database_path=mapping_database,
        )

        assert [analysis.path for analysis in analyses] == [
            unrecognized_path,
            recognized_path,
            unsupported_path,
        ]
        by_name = {analysis.path.name: analysis for analysis in analyses}
        assert by_name["Rozpoznany.xlsx"].recognized is True
        assert by_name["Rozpoznany.xlsx"].rows_count == 3
        assert by_name["Nieznany.xlsx"].recognized is False
        assert by_name["Załącznik.pdf"].recognized is False
        assert "nie jest obsługiwany" in by_name["Załącznik.pdf"].message

    def test_analysis_rejects_report_from_wrong_year(
        self,
        tmp_path,
        mapping_database,
    ):
        report_path = _write_legimi_report(tmp_path / "Legimi.xlsx")
        _save_mapping_for_table(
            report_path,
            mapping_database,
            TRANSACTION_REPORT,
            {"title": 0, "isbn": 2, "transaction_date": 5, "net": 7},
            header_row=8,
        )

        analysis = analyze_report_file(
            report_path,
            expected_year=2025,
            mapping_database_path=mapping_database,
        )

        assert analysis.recognized is False
        assert analysis.content_years == (2026,)
        assert "2025" in analysis.message
        assert "2026" in analysis.message


class TestLegacyExcelFiles:

    def test_xls_file_gets_clear_message(self, tmp_path):
        legacy_file = tmp_path / "Audioteka.xls"
        legacy_file.write_bytes(b"\xd0\xcf\x11\xe0")

        with pytest.raises(XlsxReportError, match="zapisz jako .xlsx"):
            load_report_rows(legacy_file)
