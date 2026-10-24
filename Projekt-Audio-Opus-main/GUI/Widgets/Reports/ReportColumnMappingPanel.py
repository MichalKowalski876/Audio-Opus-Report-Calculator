from pathlib import Path

from openpyxl.utils import get_column_letter
from PySide6 import QtCore, QtGui, QtWidgets

from platforms.column_mappings import (
    DEFAULT_DATABASE_PATH,
    ColumnMappingError,
    normalize_column_name,
    save_column_mapping,
)
from platforms.xlsx_reports import (
    MONTHLY_REPORT,
    QUARTERLY_REPORT,
    TRANSACTION_REPORT,
    XlsxReportError,
    find_saved_report_mapping,
    load_workbook_preview,
    report_platform_key,
    resolve_report_column_names,
)


class ReportColumnMappingPanel(QtWidgets.QWidget):
    back_requested = QtCore.Signal()
    cell_pick_requested = QtCore.Signal(str)
    mapping_saved = QtCore.Signal(str)

    FIELD_LABELS = (
        ("title", "Nazwa książki", True),
        ("isbn", "ISBN", True),
        ("author", "Autor", False),
        ("date_or_period", "Data / okres", True),
        ("quantity", "Liczba jednostek", False),
        ("net", "Wynagrodzenie netto", True),
        ("gross", "Wynagrodzenie brutto", False),
        ("vat", "VAT", False),
        ("sale_type", "Typ sprzedaży", False),
    )

    def __init__(self, parent=None, database_path=DEFAULT_DATABASE_PATH):
        super().__init__(parent)
        self.database_path = Path(database_path)
        self.report_path: Path | None = None
        self.selected_header_sheet: str | None = None
        self.selected_header_row: int | None = None
        self.selected_column_indexes: dict[str, int] = {}
        self.selected_cell_locations: dict[str, tuple[str, int]] = {}

        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(14)

        header_layout = QtWidgets.QHBoxLayout()
        self.back_button = QtWidgets.QPushButton("Wstecz")
        self.back_button.setIcon(
            self.style().standardIcon(QtWidgets.QStyle.StandardPixmap.SP_ArrowBack)
        )
        self.back_button.setMinimumWidth(110)
        self.back_button.clicked.connect(self.back_requested.emit)
        header_layout.addWidget(self.back_button)
        header_layout.addSpacing(8)
        title = QtWidgets.QLabel("Konfiguracja kolumn raportu")
        title.setProperty("role", "Title")
        header_layout.addWidget(title)
        header_layout.addStretch()
        layout.addLayout(header_layout)

        description = QtWidgets.QLabel(
            "Wpisz dokładną nazwę nagłówka z raportu albo użyj ikony Excela "
            "i wskaż komórkę zawierającą tę nazwę."
        )
        description.setWordWrap(True)
        layout.addWidget(description)

        self.report_info_label = QtWidgets.QLabel("")
        self.report_info_label.setTextInteractionFlags(
            QtCore.Qt.TextInteractionFlag.TextSelectableByMouse
        )
        layout.addWidget(self.report_info_label)

        form_group = QtWidgets.QGroupBox("Nazwy kolumn w pliku Excel")
        form = QtWidgets.QFormLayout(form_group)
        form.setHorizontalSpacing(16)
        form.setVerticalSpacing(10)

        self.report_type_combo = QtWidgets.QComboBox()
        self.report_type_combo.addItem("Pojedyncze transakcje", TRANSACTION_REPORT)
        self.report_type_combo.addItem("Raport miesięczny", MONTHLY_REPORT)
        self.report_type_combo.addItem("Raport kwartalny", QUARTERLY_REPORT)
        form.addRow("Rodzaj danych:", self.report_type_combo)

        self.field_edits: dict[str, QtWidgets.QLineEdit] = {}
        self.excel_buttons: dict[str, QtWidgets.QPushButton] = {}
        for field_name, label, required in self.FIELD_LABELS:
            row = QtWidgets.QHBoxLayout()
            edit = QtWidgets.QLineEdit()
            edit.setPlaceholderText("Wpisz nazwę nagłówka z raportu")
            edit.textEdited.connect(
                lambda _text, name=field_name: self._clear_cell_selection(name)
            )
            self.field_edits[field_name] = edit
            row.addWidget(edit, 1)

            excel_button = QtWidgets.QPushButton("▦")
            excel_button.setFixedSize(38, 32)
            excel_button.setToolTip(
                f"Otwórz podgląd Excela i wybierz kolumnę: {label}"
            )
            excel_button.clicked.connect(
                lambda _checked=False, name=field_name: self.cell_pick_requested.emit(name)
            )
            self.excel_buttons[field_name] = excel_button
            row.addWidget(excel_button)
            suffix = " *" if required else " (opcjonalnie)"
            form.addRow(f"{label}{suffix}:", row)
        layout.addWidget(form_group)

        self.header_location_label = QtWidgets.QLabel(
            "Wiersz nagłówków zostanie odnaleziony na podstawie wpisanych nazw."
        )
        self.header_location_label.setWordWrap(True)
        layout.addWidget(self.header_location_label)

        self.status_label = QtWidgets.QLabel("")
        self.status_label.setWordWrap(True)
        layout.addWidget(self.status_label)
        layout.addStretch()

        actions = QtWidgets.QHBoxLayout()
        actions.addStretch()
        self.save_button = QtWidgets.QPushButton("Zapisz konfigurację")
        self.save_button.setMinimumWidth(180)
        self.save_button.clicked.connect(self.save_configuration)
        actions.addWidget(self.save_button)
        layout.addLayout(actions)

    def configure_report(self, report_path):
        self.report_path = Path(report_path)
        self.selected_header_sheet = None
        self.selected_header_row = None
        self.selected_column_indexes.clear()
        self.selected_cell_locations.clear()
        for edit in self.field_edits.values():
            edit.clear()
        self.report_type_combo.setCurrentIndex(0)
        platform = report_platform_key(self.report_path)
        self.report_info_label.setText(
            f"Platforma: {platform}\nPlik: {self.report_path.name}"
        )
        self.header_location_label.setText(
            "Wiersz nagłówków zostanie odnaleziony na podstawie wpisanych nazw."
        )
        self._set_status("")

        for button in self.excel_buttons.values():
            button.setIcon(QtGui.QIcon())
            button.setText("▦")
        icon = QtWidgets.QFileIconProvider().icon(QtCore.QFileInfo(str(self.report_path)))
        if not icon.isNull():
            for button in self.excel_buttons.values():
                button.setText("")
                button.setIcon(icon)
                button.setIconSize(QtCore.QSize(22, 22))

        try:
            saved = find_saved_report_mapping(
                self.report_path,
                mapping_database_path=self.database_path,
            )
        except XlsxReportError as error:
            self._set_status(str(error), error=True)
            return
        if saved is None:
            return
        mapping, match = saved
        type_index = self.report_type_combo.findData(mapping.report_type)
        self.report_type_combo.setCurrentIndex(max(0, type_index))
        for field_name, column_index in mapping.columns.items():
            ui_field = (
                "date_or_period"
                if field_name in {"transaction_date", "period"}
                else field_name
            )
            if ui_field not in self.field_edits or column_index >= len(match.headers):
                continue
            self.field_edits[ui_field].setText(match.headers[column_index])
            self.selected_column_indexes[ui_field] = column_index
            self.selected_cell_locations[ui_field] = (
                match.sheet_name,
                match.header_row,
            )
        self.selected_header_sheet = match.sheet_name
        self.selected_header_row = match.header_row
        self.header_location_label.setText(
            f"Wybrany wiersz nagłówków: {match.sheet_name}, "
            f"wiersz {match.header_row}."
        )

    def field_label(self, field_name):
        return next(
            (label for name, label, _required in self.FIELD_LABELS if name == field_name),
            field_name,
        )

    def preferred_location(self):
        return self.selected_header_sheet, self.selected_header_row

    def _clear_cell_selection(self, field_name):
        self.selected_column_indexes.pop(field_name, None)
        self.selected_cell_locations.pop(field_name, None)
        if self.selected_cell_locations:
            (
                self.selected_header_sheet,
                self.selected_header_row,
            ) = next(reversed(self.selected_cell_locations.values()))
            self.header_location_label.setText(
                "Ostatnio wskazana komórka: "
                f"{self.selected_header_sheet}, "
                f"wiersz {self.selected_header_row}."
            )
        else:
            self.selected_header_sheet = None
            self.selected_header_row = None
            self.header_location_label.setText(
                "Wiersz nagłówków zostanie odnaleziony na podstawie "
                "wpisanych nazw."
            )

    def set_field_from_cell(
        self,
        field_name,
        value,
        sheet_name,
        row_number,
        column_number,
    ):
        if field_name not in self.field_edits:
            return False
        self.selected_header_sheet = sheet_name
        self.selected_header_row = int(row_number)
        self.selected_column_indexes[field_name] = int(column_number) - 1
        self.selected_cell_locations[field_name] = (
            sheet_name,
            int(row_number),
        )
        self.field_edits[field_name].setText(str(value).strip())
        self.header_location_label.setText(
            f"Ostatnio wskazana komórka: {sheet_name}, wiersz {row_number}."
        )
        self._set_status("")
        return True

    def _set_status(self, text, error=False):
        self.status_label.setText(text)
        self.status_label.setStyleSheet(
            "color: #c62828;" if error and text else ""
        )

    def save_configuration(self):
        if self.report_path is None:
            return
        required_fields = {
            field_name: label
            for field_name, label, required in self.FIELD_LABELS
            if required
        }
        missing = [
            label
            for field_name, label in required_fields.items()
            if not self.field_edits[field_name].text().strip()
        ]
        if missing:
            self._set_status(
                "Uzupełnij wymagane pola: " + ", ".join(missing) + ".",
                error=True,
            )
            return

        report_type = self.report_type_combo.currentData()
        date_field = (
            "transaction_date" if report_type == TRANSACTION_REPORT else "period"
        )
        column_names = {
            field_name: edit.text().strip()
            for field_name, edit in self.field_edits.items()
            if field_name != "date_or_period" and edit.text().strip()
        }
        column_names[date_field] = self.field_edits["date_or_period"].text().strip()
        preferred_columns = {
            (
                date_field if field_name == "date_or_period" else field_name
            ): column_index
            for field_name, column_index in self.selected_column_indexes.items()
            if (
                date_field if field_name == "date_or_period" else field_name
            ) in column_names
        }
        normalized_names = [normalize_column_name(value) for value in column_names.values()]
        if len(set(normalized_names)) != len(normalized_names):
            self._set_status(
                "Każde pole musi wskazywać inną nazwę kolumny.",
                error=True,
            )
            return

        selected_locations = {
            location
            for field_name, location in self.selected_cell_locations.items()
            if self.field_edits[field_name].text().strip()
        }
        if len(selected_locations) > 1:
            location_names = ", ".join(
                f"{sheet_name}, wiersz {row_number}"
                for sheet_name, row_number in sorted(selected_locations)
            )
            self._set_status(
                "Wszystkie nazwy kolumn wybierz z jednego wiersza nagłówków. "
                f"Wybrane komórki pochodzą z: {location_names}. "
                "Popraw wybór i zapisz konfigurację ponownie.",
                error=True,
            )
            return
        if selected_locations:
            self.selected_header_sheet, self.selected_header_row = next(
                iter(selected_locations)
            )

        try:
            match = resolve_report_column_names(
                self.report_path,
                column_names,
                preferred_sheet=self.selected_header_sheet,
                preferred_header_row=self.selected_header_row,
                preferred_columns=preferred_columns,
            )
            save_column_mapping(
                report_platform_key(self.report_path),
                match.headers,
                report_type,
                match.columns,
                database_path=self.database_path,
            )
        except (XlsxReportError, ColumnMappingError) as error:
            self._set_status(str(error), error=True)
            return

        self._set_status("Konfiguracja została zapisana.")
        self.mapping_saved.emit(str(self.report_path))


class ExcelCellPickerPanel(QtWidgets.QWidget):
    back_requested = QtCore.Signal()
    cell_selected = QtCore.Signal(str, str, int, int)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.report_path: Path | None = None
        self.previews = []

        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(12)

        header_layout = QtWidgets.QHBoxLayout()
        self.back_button = QtWidgets.QPushButton("Wstecz")
        self.back_button.setIcon(
            self.style().standardIcon(QtWidgets.QStyle.StandardPixmap.SP_ArrowBack)
        )
        self.back_button.setMinimumWidth(110)
        self.back_button.clicked.connect(self.back_requested.emit)
        header_layout.addWidget(self.back_button)
        header_layout.addSpacing(8)
        title = QtWidgets.QLabel("Wybierz nagłówek z raportu Excel")
        title.setProperty("role", "Title")
        header_layout.addWidget(title)
        header_layout.addStretch()
        layout.addLayout(header_layout)

        self.instruction_label = QtWidgets.QLabel("")
        self.instruction_label.setWordWrap(True)
        layout.addWidget(self.instruction_label)

        sheet_layout = QtWidgets.QHBoxLayout()
        sheet_layout.addWidget(QtWidgets.QLabel("Arkusz:"))
        self.sheet_combo = QtWidgets.QComboBox()
        self.sheet_combo.setMinimumWidth(240)
        self.sheet_combo.currentIndexChanged.connect(self.render_sheet)
        sheet_layout.addWidget(self.sheet_combo)
        sheet_layout.addStretch()
        self.preview_status_label = QtWidgets.QLabel("")
        sheet_layout.addWidget(self.preview_status_label)
        layout.addLayout(sheet_layout)

        self.table = QtWidgets.QTableWidget(0, 0)
        self.table.setEditTriggers(
            QtWidgets.QAbstractItemView.EditTrigger.NoEditTriggers
        )
        self.table.setSelectionMode(
            QtWidgets.QAbstractItemView.SelectionMode.SingleSelection
        )
        self.table.setSelectionBehavior(
            QtWidgets.QAbstractItemView.SelectionBehavior.SelectItems
        )
        self.table.setVerticalScrollMode(
            QtWidgets.QAbstractItemView.ScrollMode.ScrollPerPixel
        )
        self.table.setHorizontalScrollMode(
            QtWidgets.QAbstractItemView.ScrollMode.ScrollPerPixel
        )
        self.table.setWordWrap(False)
        horizontal_header = self.table.horizontalHeader()
        horizontal_header.setMinimumSectionSize(40)
        horizontal_header.setResizeContentsPrecision(-1)
        self.table.cellClicked.connect(self._cell_clicked)
        layout.addWidget(self.table, 1)

    def open_report(
        self,
        report_path,
        field_label,
        preferred_sheet=None,
        preferred_row=None,
    ):
        self.report_path = Path(report_path)
        self.instruction_label.setText(
            f"Wybierz komórkę zawierającą nazwę kolumny „{field_label}”. "
            "Po kliknięciu wartość zostanie automatycznie wpisana do konfiguratora."
        )
        self.preview_status_label.setText("Wczytywanie...")
        QtWidgets.QApplication.setOverrideCursor(QtCore.Qt.CursorShape.WaitCursor)
        try:
            self.previews = load_workbook_preview(self.report_path)
        except XlsxReportError as error:
            self.previews = []
            self.preview_status_label.setText(str(error))
        finally:
            QtWidgets.QApplication.restoreOverrideCursor()

        self.sheet_combo.blockSignals(True)
        self.sheet_combo.clear()
        for index, preview in enumerate(self.previews):
            self.sheet_combo.addItem(preview.name, index)
        if preferred_sheet:
            preferred_index = self.sheet_combo.findText(preferred_sheet)
            if preferred_index >= 0:
                self.sheet_combo.setCurrentIndex(preferred_index)
        self.sheet_combo.blockSignals(False)
        self.render_sheet()
        if preferred_row and self.table.rowCount() >= preferred_row:
            item = self.table.item(preferred_row - 1, 0)
            if item is not None:
                self.table.scrollToItem(
                    item,
                    QtWidgets.QAbstractItemView.ScrollHint.PositionAtCenter,
                )

    def render_sheet(self, _index=None):
        preview_index = self.sheet_combo.currentData()
        if preview_index is None or not self.previews:
            self.table.setRowCount(0)
            self.table.setColumnCount(0)
            return
        preview = self.previews[int(preview_index)]
        column_count = max((len(row) for row in preview.rows), default=0)
        self.table.setUpdatesEnabled(False)
        try:
            self.table.clearSpans()
            self.table.clear()
            self.table.setRowCount(len(preview.rows))
            self.table.setColumnCount(column_count)
            self.table.setHorizontalHeaderLabels(
                [get_column_letter(index + 1) for index in range(column_count)]
            )
            self.table.setVerticalHeaderLabels(
                [str(index + 1) for index in range(len(preview.rows))]
            )
            for row_index, row_values in enumerate(preview.rows):
                for column_index, value in enumerate(row_values):
                    self.table.setItem(
                        row_index,
                        column_index,
                        QtWidgets.QTableWidgetItem(value),
                    )
            for min_row, min_column, max_row, max_column in preview.merged_ranges:
                row_span = max_row - min_row + 1
                column_span = max_column - min_column + 1
                if row_span > 1 or column_span > 1:
                    self.table.setSpan(
                        min_row - 1,
                        min_column - 1,
                        row_span,
                        column_span,
                    )
        finally:
            self.table.setUpdatesEnabled(True)
        self.table.resizeColumnsToContents()
        status = f"{preview.total_rows} wierszy, {preview.total_columns} kolumn"
        if preview.truncated:
            status += " — podgląd ograniczony"
        self.preview_status_label.setText(status)

    def _cell_clicked(self, row, column):
        item = self.table.item(row, column)
        value = item.text().strip() if item is not None else ""
        if not value:
            self.preview_status_label.setText("Wybrana komórka jest pusta.")
            return
        sheet_name = self.sheet_combo.currentText()
        self.cell_selected.emit(value, sheet_name, row + 1, column + 1)
