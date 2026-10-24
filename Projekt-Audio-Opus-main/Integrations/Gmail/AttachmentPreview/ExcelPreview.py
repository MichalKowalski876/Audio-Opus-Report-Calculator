from io import BytesIO

from PySide6 import QtWidgets, QtCore
from openpyxl import load_workbook


class ExcelPreview(QtWidgets.QWidget):
    def __init__(self, data, parent=None):
        super().__init__(parent)

        self.data = data

        self.create_ui()

    def create_ui(self):
        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)

        try:
            workbook = load_workbook(
                filename=BytesIO(self.data),
                read_only=True,
                data_only=True
            )

            tabs = QtWidgets.QTabWidget()

            for worksheet in workbook.worksheets:
                table = self.create_sheet_table(worksheet)

                tabs.addTab(
                    table,
                    worksheet.title
                )

            layout.addWidget(tabs)

            workbook.close()

        except Exception as error:
            error_label = QtWidgets.QLabel(
                "Nie udało się wyświetlić pliku Excel.\n\n"
                f"{error}"
            )

            error_label.setWordWrap(True)

            layout.addWidget(error_label)

    def create_sheet_table(self, worksheet):
        table = QtWidgets.QTableWidget()

        rows = list(worksheet.iter_rows(values_only=True))

        if not rows:
            table.setRowCount(0)
            table.setColumnCount(0)
            return table

        column_count = max(
            len(row) for row in rows
        )

        table.setRowCount(len(rows))
        table.setColumnCount(column_count)

        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                if value is None:
                    text = ""
                else:
                    text = str(value)

                item = QtWidgets.QTableWidgetItem(text)

                item.setFlags(
                    item.flags()
                    & ~QtCore.Qt.ItemFlag.ItemIsEditable
                )

                table.setItem(
                    row_index,
                    column_index,
                    item
                )

        table.resizeColumnsToContents()

        return table