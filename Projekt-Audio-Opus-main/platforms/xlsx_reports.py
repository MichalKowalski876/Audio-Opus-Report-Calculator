"""Ogólne rozpoznawanie raportów platform zapisanych w plikach Excel.

Parser nie zakłada nazw ani układu kolumn żadnej platformy. Korzysta wyłącznie
z mapowań zapisanych przez użytkownika dla platformy i układu Excela.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .column_mappings import (
    DEFAULT_DATABASE_PATH,
    ColumnMappingError,
    SavedColumnMapping,
    layout_signature,
    load_platform_column_mappings,
    normalize_column_name,
)
from .products import find_product, normalize_isbn, normalize_title


TRANSACTION_REPORT = "transaction"
MONTHLY_REPORT = "monthly"
QUARTERLY_REPORT = "quarterly"

SUPPORTED_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
LEGACY_EXCEL_SUFFIXES = {".xls", ".xlsb", ".xlm"}
IGNORED_REPORT_SUFFIXES = {".json"}

KNOWN_PLATFORM_KEYS = ("legimi", "storytel", "audioteka", "voices")
SUMMARY_TITLES = {"razem", "suma", "total", "sum", "subtotal", "sum total"}


class XlsxReportError(RuntimeError):
    pass


@dataclass(frozen=True)
class ImportedRow:
    """Pojedyncza pozycja rozpoznana w dowolnym raporcie Excel."""

    title: str
    isbn: str = ""
    net: float = 0.0
    gross: float | None = None
    month: str = ""
    transaction_date: str = ""
    period_start: str = ""
    period_end: str = ""
    quantity: float | None = None
    report_type: str = MONTHLY_REPORT
    # Autor z raportu (np. kolumna "Author(s)" u Storytel). Rozstrzyga, gdy
    # platforma poda ten sam ISBN albo ten sam tytuł dla dwóch różnych książek.
    author: str = ""
    author: str = ""


@dataclass
class MatchedProduct:
    product: dict
    net: float = 0.0
    gross: float | None = None
    rows: int = 0


@dataclass
class UnmatchedTitle:
    title: str
    net: float = 0.0
    gross: float | None = None
    rows: int = 0
    isbns: set = field(default_factory=set)
    author: str = ""


@dataclass(frozen=True)
class ReportFileAnalysis:
    path: Path
    recognized: bool
    platform_key: str = ""
    report_types: tuple[str, ...] = ()
    content_years: tuple[int, ...] = ()
    rows_count: int = 0
    message: str = ""


@dataclass(frozen=True)
class ColumnMappingMatch:
    sheet_name: str
    header_row: int
    headers: tuple[str, ...]
    columns: dict[str, int]


@dataclass(frozen=True)
class WorkbookSheetPreview:
    name: str
    rows: tuple[tuple[str, ...], ...]
    total_rows: int
    total_columns: int
    merged_ranges: tuple[tuple[int, int, int, int], ...] = ()
    truncated: bool = False


def normalize_header(value: Any) -> str:
    return normalize_column_name(value)


CURRENCY_PATTERN = re.compile(r"(?:zl|zł|pln|eur|usd|€|\$)", re.IGNORECASE)


def _to_float(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"\s", "", str(value))
    text = CURRENCY_PATTERN.sub("", text)
    if not text:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _normalize_isbn_value(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text


def _is_real_isbn(value) -> bool:
    return bool(_normalize_isbn_value(value))


def _find_header_rows(
    sheet,
    saved_mappings: dict[str, SavedColumnMapping] | None = None,
) -> list[int]:
    saved_signatures = set((saved_mappings or {}).keys())
    return [
        row_number
        for row_number, row_values in enumerate(
            sheet.iter_rows(values_only=True), start=1
        )
        if layout_signature(row_values) in saved_signatures
    ]


def infer_period_from_text(value: Any) -> str:
    text = " ".join(str(value or "").split())
    if not text:
        return ""
    year_first = re.search(
        r"(?<!\d)(20\d{2})[-_. /](0?[1-9]|1[0-2])(?:[-_. /]\d{1,2})?(?!\d)",
        text,
    )
    if year_first:
        return f"{int(year_first.group(1)):04d}-{int(year_first.group(2)):02d}"
    month_first = re.search(
        r"(?<!\d)(0?[1-9]|1[0-2])[-_. /](20\d{2})(?!\d)",
        text,
    )
    if month_first:
        return f"{int(month_first.group(2)):04d}-{int(month_first.group(1)):02d}"
    return ""


def _month_from_value(value) -> str:
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}-{value.month:02d}"

    period = infer_period_from_text(value)
    if period:
        return period

    start_date, _ = _date_range_from_value(value)
    if start_date and len(start_date) >= 7:
        return start_date[:7]

    single_date = _date_from_value(value)
    if single_date and len(single_date) >= 7:
        return single_date[:7]

    return ""



def _date_from_value(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = " ".join(str(value or "").split())
    year_first = re.search(
        r"(?<!\d)(\d{4})[-./](\d{1,2})(?:[-./](\d{1,2}))?(?!\d)", text
    )
    if year_first:
        year, month, day = year_first.groups()
        try:
            return date(int(year), int(month), int(day or 1)).isoformat()
        except ValueError:
            return ""
    day_first = re.search(
        r"(?<!\d)(\d{1,2})[-./](\d{1,2})[-./](\d{4})(?!\d)", text
    )
    if day_first:
        day, month, year = day_first.groups()
        try:
            return date(int(year), int(month), int(day)).isoformat()
        except ValueError:
            return ""
    return ""


def _date_range_from_value(value) -> tuple[str, str]:
    """Odczytuje pierwszą i ostatnią datę z zakresu rozliczeniowego."""
    if isinstance(value, datetime):
        result = value.date().isoformat()
        return result, result
    if isinstance(value, date):
        result = value.isoformat()
        return result, result

    text = " ".join(str(value or "").split())
    parsed = []
    for match in re.finditer(
        r"(?<!\d)(\d{4})[-./](\d{1,2})[-./](\d{1,2})(?!\d)",
        text,
    ):
        try:
            parsed.append(
                date(*(int(part) for part in match.groups())).isoformat()
            )
        except ValueError:
            continue
    if not parsed:
        for match in re.finditer(
            r"(?<!\d)(\d{1,2})[-./](\d{1,2})[-./](\d{4})(?!\d)",
            text,
        ):
            day, month, year = match.groups()
            try:
                parsed.append(
                    date(int(year), int(month), int(day)).isoformat()
                )
            except ValueError:
                continue
    if not parsed:
        return "", ""
    return parsed[0], parsed[-1]


AUTHOR_HEADERS = {"author", "authors", "author s", "autor", "autorzy", "autor ka"}


def _author_column(header_values) -> int | None:
    """Kolumna autora rozpoznana po nagłówku ("Author(s)", "Autor", ...)."""
    for index, value in enumerate(header_values or ()):
        if normalize_header(value) in AUTHOR_HEADERS:
            return index
    return None


def _cell(row_values, mapping: dict[str, int], field_name: str):
    index = mapping.get(field_name)
    if index is None or index >= len(row_values):
        return None
    return row_values[index]


def _read_table(
    file_path: Path,
    sheet,
    header_row: int,
    end_row: int,
    saved_mappings: dict[str, SavedColumnMapping] | None = None,
) -> tuple[list[ImportedRow], str]:
    header_values = next(
        sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True),
        None,
    )
    if not header_values:
        return [], "pusty wiersz nagłówka"

    saved_mapping = (saved_mappings or {}).get(layout_signature(header_values))
    if saved_mapping is None:
        return [], "brak zapisanego mapowania kolumn"
    mapping = dict(saved_mapping.columns)
    missing = [field for field in ("title", "isbn", "net") if field not in mapping]
    if missing:
        return [], "brak wymaganych kolumn: " + ", ".join(missing)

    report_type = saved_mapping.report_type
    if not report_type:
        return [], "nie rozpoznano typu raportu"

    # Mapowania zapisane przed dodaniem pola "author" go nie mają - wtedy
    # szukamy kolumny autora po nagłówku, żeby nie trzeba było mapować
    # raportu od nowa.
    if "author" not in mapping:
        author_column = _author_column(header_values)
        if author_column is not None:
            mapping["author"] = author_column

    inferred_period = infer_period_from_text(file_path.stem) or infer_period_from_text(
        sheet.title
    )
    rows = []
    for row_values in sheet.iter_rows(
        min_row=header_row + 1,
        max_row=end_row,
        values_only=True,
    ):
        title = str(_cell(row_values, mapping, "title") or "").strip()
        if not title or normalize_title(title).strip(" .:-") in SUMMARY_TITLES:
            continue
        isbn_value = _cell(row_values, mapping, "isbn")
        if not _is_real_isbn(isbn_value):
            continue
        net = _to_float(_cell(row_values, mapping, "net"))
        if net is None:
            continue
        gross = _to_float(_cell(row_values, mapping, "gross"))
        if gross is None:
            vat = _to_float(_cell(row_values, mapping, "vat"))
            if vat is not None:
                gross = net * (1 + vat / 100)
        period_value = _cell(
            row_values,
            mapping,
            "transaction_date" if report_type == TRANSACTION_REPORT else "period",
        )
        month = _month_from_value(period_value) or inferred_period
        period_start, period_end = _date_range_from_value(period_value)
        transaction_date = (
            _date_from_value(period_value)
            if report_type == TRANSACTION_REPORT
            else ""
        )
        quantity = _to_float(_cell(row_values, mapping, "quantity"))
        author_value = _cell(row_values, mapping, "author")
        author = " ".join(str(author_value or "").split())
        rows.append(
            ImportedRow(
                title=title,
                isbn=_normalize_isbn_value(isbn_value),
                net=net,
                gross=gross,
                month=month,
                transaction_date=transaction_date,
                period_start=period_start,
                period_end=period_end,
                quantity=quantity,
                report_type=report_type,
                author=author,
            )
        )
    if not rows and saved_mapping is not None:
        return [], "zapamiętane mapowanie nie zawiera prawidłowych pozycji"
    return rows, ""


def detect_platform_key(path):
    """Rozpoznaje wyłącznie nazwę platformy, nie układ kolumn raportu."""
    stem = Path(path).stem.strip().casefold()
    for key in KNOWN_PLATFORM_KEYS:
        if key in stem:
            return key
    return None


def _fallback_platform_key(path: Path) -> str:
    return path.stem.strip().casefold() or "platforma"


def report_platform_key(path, platform_key=None) -> str:
    key = str(platform_key or "").strip().casefold()
    if key:
        return key
    path = Path(path)
    return detect_platform_key(path) or _fallback_platform_key(path)


def find_saved_report_mapping(
    path,
    mapping_database_path=DEFAULT_DATABASE_PATH,
) -> tuple[SavedColumnMapping, ColumnMappingMatch] | None:
    """Zwraca pierwsze mapowanie zapisane dla układu obecnego w raporcie."""
    path = _excel_path(path)
    key = report_platform_key(path)
    try:
        saved_mappings = load_platform_column_mappings(
            key,
            database_path=mapping_database_path,
        )
    except ColumnMappingError as error:
        raise XlsxReportError(str(error)) from error
    if not saved_mappings:
        return None

    try:
        workbook = load_workbook(path, data_only=True, read_only=False)
    except Exception as error:
        raise XlsxReportError(f"Nie udało się otworzyć pliku: {error}") from error
    try:
        for sheet in workbook.worksheets:
            for row_number, row_values in enumerate(
                sheet.iter_rows(values_only=True),
                start=1,
            ):
                saved_mapping = saved_mappings.get(layout_signature(row_values))
                if saved_mapping is None:
                    continue
                headers = tuple(
                    "" if value is None else str(value).strip()
                    for value in row_values
                )
                return saved_mapping, ColumnMappingMatch(
                    sheet_name=sheet.title,
                    header_row=row_number,
                    headers=headers,
                    columns=dict(saved_mapping.columns),
                )
    finally:
        workbook.close()
    return None


def load_report_rows(
    path,
    platform_key=None,
    mapping_database_path=DEFAULT_DATABASE_PATH,
):
    """Czyta wszystkie rozpoznane tabele i zwraca wspólne pozycje raportu."""
    path = Path(path)
    if not path.exists():
        raise XlsxReportError(f"Plik nie istnieje: {path}")
    if path.suffix.lower() in LEGACY_EXCEL_SUFFIXES:
        raise XlsxReportError(
            f"Format {path.suffix.lower()} to stary plik Excela, którego program "
            "nie czyta. Otwórz go w Excelu i zapisz jako .xlsx."
        )
    if path.suffix.lower() not in SUPPORTED_EXCEL_SUFFIXES:
        raise XlsxReportError(f"Format {path.suffix.lower()} nie jest obsługiwany.")
    key = report_platform_key(path, platform_key)
    try:
        saved_mappings = load_platform_column_mappings(
            key,
            database_path=mapping_database_path,
        )
    except ColumnMappingError as error:
        raise XlsxReportError(str(error)) from error

    try:
        workbook = load_workbook(path, data_only=True, read_only=False)
    except Exception as error:
        raise XlsxReportError(f"Nie udało się otworzyć pliku: {error}") from error

    rows = []
    problems = []
    try:
        for sheet in workbook.worksheets:
            header_rows = _find_header_rows(sheet, saved_mappings)
            if not header_rows:
                continue
            for index, header_row in enumerate(header_rows):
                end_row = (
                    header_rows[index + 1] - 1
                    if index + 1 < len(header_rows)
                    else (sheet.max_row or header_row)
                )
                table_rows, problem = _read_table(
                    path,
                    sheet,
                    header_row,
                    end_row,
                    saved_mappings,
                )
                rows.extend(table_rows)
                if problem:
                    problems.append(f"{sheet.title}, wiersz {header_row}: {problem}")
    finally:
        workbook.close()

    if not rows:
        detail = "; ".join(problems)
        if not detail:
            detail = (
                f"brak zapisanego mapowania kolumn dla platformy {key}. "
                "Użyj przycisku Dopasuj kolumny"
            )
        raise XlsxReportError(f"Nie rozpoznano raportu: {detail}.")

    return key, rows


def _excel_path(path) -> Path:
    path = Path(path)
    if not path.exists():
        raise XlsxReportError(f"Plik nie istnieje: {path}")
    if path.suffix.lower() not in SUPPORTED_EXCEL_SUFFIXES:
        raise XlsxReportError(
            "Konfiguracja kolumn jest dostępna tylko dla raportów Excel "
            "w formacie .xlsx."
        )
    return path


def resolve_report_column_names(
    path,
    column_names: dict[str, str],
    preferred_sheet: str | None = None,
    preferred_header_row: int | None = None,
    preferred_columns: dict[str, int] | None = None,
) -> ColumnMappingMatch:
    """Znajduje wiersz zawierający dokładnie skonfigurowane nagłówki."""
    path = _excel_path(path)
    expected = {
        field_name: normalize_header(column_name)
        for field_name, column_name in column_names.items()
        if normalize_header(column_name)
    }
    if not expected:
        raise XlsxReportError("Nie podano żadnej nazwy kolumny.")
    if len(set(expected.values())) != len(expected):
        raise XlsxReportError("Każde pole musi wskazywać inną nazwę kolumny.")

    try:
        workbook = load_workbook(path, data_only=True, read_only=False)
    except Exception as error:
        raise XlsxReportError(f"Nie udało się otworzyć pliku: {error}") from error

    matches: list[ColumnMappingMatch] = []
    try:
        for sheet in workbook.worksheets:
            if preferred_sheet and sheet.title != preferred_sheet:
                continue
            row_numbers = (
                [preferred_header_row]
                if preferred_header_row is not None
                else range(1, (sheet.max_row or 0) + 1)
            )
            for row_number in row_numbers:
                if row_number is None or not 1 <= row_number <= (sheet.max_row or 0):
                    continue
                raw_headers = next(
                    sheet.iter_rows(
                        min_row=row_number, max_row=row_number, values_only=True
                    ),
                    (),
                )
                positions: dict[str, list[int]] = {}
                for column_index, value in enumerate(raw_headers):
                    normalized = normalize_header(value)
                    if normalized:
                        positions.setdefault(normalized, []).append(column_index)
                columns = {}
                for field_name, column_name in expected.items():
                    selected_column = (preferred_columns or {}).get(field_name)
                    if selected_column is not None:
                        if (
                            selected_column >= len(raw_headers)
                            or normalize_header(raw_headers[selected_column])
                            != column_name
                        ):
                            columns = {}
                            break
                        columns[field_name] = selected_column
                        continue
                    matching_columns = positions.get(column_name, ())
                    if len(matching_columns) != 1:
                        columns = {}
                        break
                    columns[field_name] = matching_columns[0]
                if len(columns) != len(expected):
                    continue
                headers = tuple(
                    "" if value is None else str(value).strip()
                    for value in raw_headers
                )
                matches.append(
                    ColumnMappingMatch(
                        sheet_name=sheet.title,
                        header_row=row_number,
                        headers=headers,
                        columns=columns,
                    )
                )
    finally:
        workbook.close()

    if not matches:
        names = ", ".join(str(value) for value in column_names.values() if value)
        location = (
            f" w arkuszu {preferred_sheet}, wiersz {preferred_header_row}"
            if preferred_sheet and preferred_header_row
            else ""
        )
        raise XlsxReportError(
            f"Nie znaleziono jednego wiersza{location} zawierającego kolumny: "
            f"{names}."
        )
    return matches[0]


def load_workbook_preview(
    path,
    maximum_rows: int | None = None,
    maximum_columns: int | None = None,
) -> list[WorkbookSheetPreview]:
    """Ładuje wypełniony obszar każdego arkusza do podglądu w GUI."""
    path = _excel_path(path)
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as error:
        raise XlsxReportError(f"Nie udało się otworzyć pliku: {error}") from error

    previews: list[WorkbookSheetPreview] = []
    try:
        for sheet in workbook.worksheets:
            loaded_rows: list[tuple[str, ...]] = []
            total_rows = 0
            total_columns = 0
            for row_index, cells in enumerate(sheet.iter_rows(), start=1):
                row_values = tuple(
                    "" if cell.value is None else str(cell.value)
                    for cell in cells
                )
                loaded_rows.append(row_values)
                non_empty_columns = [
                    index
                    for index, value in enumerate(row_values, start=1)
                    if value != ""
                ]
                if non_empty_columns:
                    total_rows = row_index
                    total_columns = max(total_columns, non_empty_columns[-1])

            visible_rows = total_rows
            if maximum_rows is not None:
                visible_rows = min(total_rows, max(1, int(maximum_rows)))
            visible_columns = total_columns
            if maximum_columns is not None:
                visible_columns = min(total_columns, max(1, int(maximum_columns)))

            rows = tuple(
                row_values[:visible_columns]
                for row_values in loaded_rows[:visible_rows]
            )
            merged_ranges = []
            for merged_range in sheet.merged_cells.ranges:
                if (
                    merged_range.min_row > visible_rows
                    or merged_range.min_col > visible_columns
                ):
                    continue
                merged_ranges.append(
                    (
                        merged_range.min_row,
                        merged_range.min_col,
                        min(merged_range.max_row, visible_rows),
                        min(merged_range.max_col, visible_columns),
                    )
                )
            previews.append(
                WorkbookSheetPreview(
                    name=sheet.title,
                    rows=rows,
                    total_rows=total_rows,
                    total_columns=total_columns,
                    merged_ranges=tuple(merged_ranges),
                    truncated=(
                        visible_rows < total_rows
                        or visible_columns < total_columns
                    ),
                )
            )
    finally:
        workbook.close()
    return previews


def rows_for_monthly_summary(rows: Iterable[ImportedRow]) -> list[ImportedRow]:
    """Nie dubluje okresu, gdy plik zawiera transakcje i ich podsumowanie."""
    rows = list(rows)
    monthly_periods = {
        row.month
        for row in rows
        if row.report_type == MONTHLY_REPORT and row.month
    }
    return [
        row
        for row in rows
        if row.report_type == MONTHLY_REPORT
        or (
            row.report_type == TRANSACTION_REPORT
            and row.month not in monthly_periods
        )
    ]


def report_content_years(rows: Iterable[ImportedRow]) -> tuple[int, ...]:
    """Zwraca lata rzeczywiście występujące w rozpoznanych pozycjach raportu."""
    years = set()
    for row in rows:
        for value in (
            row.month,
            row.transaction_date,
            row.period_start,
            row.period_end,
        ):
            text = str(value or "").strip()
            if len(text) >= 4 and text[:4].isdigit():
                years.add(int(text[:4]))
    return tuple(sorted(years))


def detect_report_years(
    path,
    platform_key=None,
    mapping_database_path=DEFAULT_DATABASE_PATH,
) -> tuple[int, ...]:
    """Rozpoznaje lata z danych, korzystając z zapisanego mapowania kolumn."""
    _key, rows = load_report_rows(
        path,
        platform_key=platform_key,
        mapping_database_path=mapping_database_path,
    )
    return report_content_years(rows)


def analyze_report_file(
        path,
        expected_year: int | None = None,
        mapping_database_path=DEFAULT_DATABASE_PATH,
) -> ReportFileAnalysis:
    path = Path(path)
    try:
        platform_key, rows = load_report_rows(
            path,
            mapping_database_path=mapping_database_path,
        )
    except XlsxReportError as error:
        return ReportFileAnalysis(path=path, recognized=False, message=str(error))

    report_types = tuple(sorted({row.report_type for row in rows}))
    content_years = report_content_years(rows)
    missing_periods = sum(
        (
            not row.period_start or not row.period_end
            if row.report_type == QUARTERLY_REPORT
            else not row.month
        )
        for row in rows
    )
    if missing_periods:
        return ReportFileAnalysis(
            path=path,
            recognized=False,
            platform_key=platform_key,
            report_types=report_types,
            content_years=content_years,
            rows_count=len(rows),
            message=f"Nie rozpoznano okresu w {missing_periods} pozycjach.",
        )

    if expected_year is not None and expected_year not in content_years:
        if content_years:
            years_text = ", ".join(str(year) for year in content_years)
            message = (
                f"Raport zawiera dane z lat ({years_text}), ale brak w nim "
                f"pozycji z wybranego roku {expected_year}."
            )
        else:
            message = f"Raport nie zawiera pozycji z roku {expected_year}."
        return ReportFileAnalysis(
            path=path,
            recognized=False,
            platform_key=platform_key,
            report_types=report_types,
            content_years=content_years,
            rows_count=len(rows),
            message=message,
        )

    return ReportFileAnalysis(
        path=path,
        recognized=True,
        platform_key=platform_key,
        report_types=report_types,
        content_years=content_years,
        rows_count=len(rows),
        message=f"Rozpoznano {len(rows)} pozycji.",
    )

def scan_report_directory(
    directory,
    mapping_database_path=DEFAULT_DATABASE_PATH,
) -> list[ReportFileAnalysis]:
    directory = Path(directory)
    if not directory.is_dir():
        return []
    files = sorted(
        (
            path
            for path in directory.rglob("*")
            if path.is_file()
            and not path.name.startswith((".", "~$"))
            and not any(
                part.startswith(".")
                for part in path.relative_to(directory).parts[:-1]
            )
            and path.suffix.lower() not in IGNORED_REPORT_SUFFIXES
        ),
        key=lambda path: str(path).casefold(),
    )
    analyses = []
    for path in files:
        relative_path = path.relative_to(directory)
        expected_year = next(
            (
                int(part)
                for part in relative_path.parts[:-1]
                if part.isdigit() and len(part) == 4
            ),
            None,
        )
        analyses.append(
            analyze_report_file(
                path,
                expected_year=expected_year,
                mapping_database_path=mapping_database_path,
            )
        )
    return analyses


def _products_indexes(products):
    by_isbn = {}
    by_title = {}
    for product in products:
        if not isinstance(product, dict) or "id" not in product:
            continue
        isbn = normalize_isbn(product.get("isbn"))
        if isbn:
            by_isbn.setdefault(isbn, product)
        title = normalize_title(
            product.get("title") or product.get("tytuł") or product.get("tytul")
        )
        if title:
            by_title.setdefault(title, product)
    return by_isbn, by_title


def match_rows_to_products(rows, products):
    """Sumuje rozpoznane pozycje po ISBN, a awaryjnie po tytule książki."""
    by_isbn, by_title = _products_indexes(products)
    matched = {}
    unmatched = {}
    for row in rows:
        isbn = normalize_isbn(row.isbn)
        title = normalize_title(row.title)
        product = find_product(row.title, row.isbn, row.author, by_isbn, by_title)
        if product is not None:
            key = product["id"]
            summary = matched.setdefault(key, MatchedProduct(product=product))
        else:
            key = title or isbn or row.title.casefold()
            summary = unmatched.setdefault(key, UnmatchedTitle(title=row.title))
            if row.isbn:
                summary.isbns.add(row.isbn)
            if row.author and not summary.author:
                summary.author = row.author
        summary.net += row.net
        if row.gross is not None:
            summary.gross = (summary.gross or 0.0) + row.gross
        summary.rows += 1

    for summary in (*matched.values(), *unmatched.values()):
        summary.net = round(summary.net, 2)
        if summary.gross is not None:
            summary.gross = round(summary.gross, 2)
    return list(matched.values()), list(unmatched.values())
